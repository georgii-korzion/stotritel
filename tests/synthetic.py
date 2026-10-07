"""Синтетический прайс в формате Fimex (раздел 4 ТЗ) — для тестов без образца и без сети.

Строка названия и пустой разделитель объединены A:E, как в настоящей выгрузке;
цены от 1000 записываются строкой с пробелом ('1 695').
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import Workbook

Row = tuple  # (артикул, ETA, регион, цена, количество)
Block = tuple[str, list[Row]]


def fimex_price(value):
    """Как Fimex: число, если меньше 1000; иначе строка с пробелом-разделителем."""
    if isinstance(value, str) or value < 1000:
        return value
    return f"{value:,}".replace(",", " ")


def write_pricelist(path: Path, sheets: dict[str, Iterable[Block | tuple]]) -> Path:
    """sheets: {бренд: [(название, [строки предложений]), ..., ("raw", (a, b, c, d, e))]}."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, blocks in sheets.items():
        ws = wb.create_sheet(name)
        r = 1
        for block in blocks:
            if block[0] == "raw":
                for c, v in enumerate(block[1], start=1):
                    if v is not None:
                        ws.cell(r, c, v)
                r += 1
                continue
            title, rows = block
            ws.cell(r, 1, title)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
            r += 1
            for article, eta, region, price, qty in rows:
                for c, v in enumerate((article, eta, region, fimex_price(price), qty), start=1):
                    ws.cell(r, c, v)
                r += 1
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)  # разделитель
            r += 1
    wb.save(path)
    return path


# Небольшой прайс со всеми особенностями из раздела 4.
SMALL: dict[str, list[Block]] = {
    "Apple": [
        ("iPhone 15 128Gb Black", [
            ("MTLV3 LL/A", "08 Oct", "🇺🇸 US", 572, 9),
            ("MTP03 AA/A", "08 Oct", "🇦🇪 AE", 592, 30),
            ("MTP03 HN/A", "08 Oct", "🇮🇳 IN", 612, 50),
        ]),
        ("iPhone 15 128Gb Blue", [("MTP43 HN/A", "08 Oct", "🇮🇳 IN", 597, 50)]),
        ("iPhone 17 Pro Max 1Tb Blue", [
            ("MFYH4 J/A", "08 Oct", "🇯🇵 JP", 1494, 19),
            ("MFYX4 X/A", "08 Oct", "🇦🇺 AU", 1805, 2),
        ]),
        ("iPad mini A17 Pro 128Gb Space Gray", [
            ("MXN73 LL/A", "08 Oct", "🇺🇸 US", 631, 5),
            ("MXN73 LL/A", "17 Oct", "🇺🇸 US", 681, 30),
        ]),
        ("MacBook Pro 14 M4 Pro 24/1Tb Silver", [
            ("MX2E3 LL/A", "08 Oct", "🇺🇸 US", "2 199", 4),  # неразрывный пробел
            ("MX2E3 ZP/A", "08 Oct", "🇸🇬 SG", 2399, 3),
        ]),
    ],
    "Samsung": [
        ("Galaxy S25 Ultra 256Gb Black", [
            (37015477980, "08 Oct", "🇦🇪 AE", 1045, 14),  # артикул — целое число
            ("07977974724", "17 Oct", "🇪🇺 GL", 1099, 6),  # артикул с ведущим нулём
        ]),
        ("Tab S10 FE 8+ 128Gb Blue", [(37015481234, "08 Oct", "🇦🇪 AE", 468, 3)]),
        ("Galaxy Z Fold7 512Gb Silver", [(37015499999, "08 Oct", "🇭🇰 HK", 1755, 5)]),
        ("Tab S10 FE 8+ 128Gb Blue", [(37015485678, "08 Oct", "🇦🇪 AE", 341, 10)]),  # второй блок
    ],
    "Meta": [
        ("Ray-Ban Meta Wayfarer Matte Black", [("RW4006", "08 Oct", "👓", 299, 12)]),
    ],
    "Sony": [
        ("DualSense Edge", [("CFI-ZCP1", "08 Oct", "🎮", 189, 7)]),
        ("PlayStation 5 Pro", [("CFI-7021", "08 Oct", "🇯🇵 JP", "1 049", 3)]),  # узкий неразрывный
    ],
    "Marshall": [
        ("Major V Black", [("1006832", "08 Oct", "🎧", 119, 25)]),
    ],
}

# Прайс побольше — для прогона этапа 3 на синтетике: есть кандидаты на все четыре правки.
DEMO: dict[str, list[Block]] = {
    "Apple": SMALL["Apple"] + [
        ("iPhone 17 Pro 256Gb Silver", [
            ("MG8G4 J/A", "08 Oct", "🇯🇵 JP", 1105, 40),
            ("MG8H4 ZA/A", "08 Oct", "🇭🇰 HK", 1139, 25),
            ("MG8J4 X/A", "08 Oct", "🇦🇺 AU", 1189, 6),
        ]),
        ("MacBook Pro 16 M4 Max 48/1Tb Space Black", [("MX313 LL/A", "08 Oct", "🇺🇸 US", 3899, 2)]),
        ("AirPods Pro 2 USB-C", [("MTJV3 AM/A", "08 Oct", "🎧", 189, 120)]),
        ("Apple Watch Ultra 2 49mm Black Ti", [
            ("MX4V3 LL/A", "08 Oct", "🇺🇸 US", 705, 11),
            ("MX4V3 AE/A", "08 Oct", "🇦🇪 AE", 722, 8),
        ]),
        ("Mac Studio M4 Max 36/512", [("MU963 LL/A", "17 Oct", "🇺🇸 US", 1999, 3)]),
        ("iPhone 17 Pro Max 2Tb Deep Blue", [
            ("MFZ04 J/A", "08 Oct", "🇯🇵 JP", 1890, 4),
            ("MFZ14 ZA/A", "08 Oct", "🇭🇰 HK", 2145, 3),
            ("MFZ24 X/A", "08 Oct", "🇦🇺 AU", 2270, 1),
        ]),
    ],
    "Samsung": SMALL["Samsung"],
    "Xiaomi": [
        ("Redmi Note 14 Pro 8/256 Black", [(41300012345, "08 Oct", "🇪🇺 GL", 249, 60)]),
        ("Xiaomi 15 Ultra 16/512 White", [
            (41300054321, "08 Oct", "🇨🇳 CN", 1129, 9),
            (41300054322, "08 Oct", "🇪🇺 GL", 1239, 4),
        ]),
    ],
    "Meta": SMALL["Meta"],
    "Sony": SMALL["Sony"],
    "Marshall": SMALL["Marshall"],
}
